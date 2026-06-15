"""
Generación de informe: consola + Excel (.xlsx).
"""

from src.accounting.fifo_calculator import TaxEngine
from src.tax_compliance.core.irpf_calculator import IRPFCalculator

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# ─── Colores corporativos ─────────────────────────────────────────────────────
COLOR_HEADER = "1F4E79"  # azul oscuro
COLOR_SUBHEAD = "2E75B6"  # azul medio
COLOR_GAIN = "C6EFCE"  # verde claro
COLOR_LOSS = "FFC7CE"  # rojo claro
COLOR_NEUTRAL = "DDEBF7"  # azul muy claro
COLOR_TOTAL = "FFF2CC"  # amarillo suave
COLOR_WHITE = "FFFFFF"


def _eur(val: float) -> str:
    return f"{val:,.2f} €"


# ─── Informe de consola ───────────────────────────────────────────────────────


def print_console_report(calc: IRPFCalculator):
    s = calc.summary()
    tax_year = s["año_fiscal"]
    line = "=" * 65

    print(f"\n{line}")
    print(f"  DECLARACIÓN DE LA RENTA {tax_year}  —  Resumen IRPF")
    print(line)

    # ── Rendimientos del trabajo
    if s["rendimiento_trabajo_bruto"] > 0:
        print("\n📋  RENDIMIENTOS DEL TRABAJO")
        print(f"  Rendimiento íntegro:           {_eur(s['rendimiento_trabajo_bruto'])}")
        print(f"  Cotización SS (trabajador):    {_eur(s['ss_trabajador'])}")
        print(f"  Base imponible general:        {_eur(s['base_general'])}")
        print(f"  Cuota íntegra (trabajo):       {_eur(s['cuota_general'])}")

    # ── Base del Ahorro — GPP
    print("\n📈  GANANCIAS Y PÉRDIDAS PATRIMONIALES (acciones + cripto)")
    print(f"  Ganancias brutas:              {_eur(s['ganancias_patrimoniales_bruto'])}")
    print(f"  Pérdidas brutas:               {_eur(s['perdidas_patrimoniales_bruto'])}")
    print(f"  GPP NETO:                      {_eur(s['gpp_neto'])}")

    # ── Base del Ahorro — RCM
    print("\n💰  RENDIMIENTOS DEL CAPITAL MOBILIARIO (dividendos)")
    print(f"  Dividendos brutos:             {_eur(s['dividendos_brutos'])}")
    print(f"  Retención extranjera:          {_eur(s['retencion_extranjera_dividendos'])}")
    print(f"  Retención española:            {_eur(s['retencion_espania_dividendos'])}")

    # ── Resumen base ahorro
    print("\n🏦  BASE IMPONIBLE DEL AHORRO")
    print(f"  Base del ahorro:               {_eur(s['base_ahorro'])}")
    print(f"  Cuota íntegra (ahorro):        {_eur(s['cuota_ahorro'])}")
    print(f"  Deducción doble imposición:  - {_eur(s['deduccion_doble_imposicion'])}")

    # ── Resultado
    print(f"\n{'─' * 65}")
    print(f"  Cuota íntegra total:           {_eur(s['cuota_integra'])}")
    print(f"  Cuota líquida:                 {_eur(s['cuota_liquida'])}")
    print(f"  Total retenciones pagadas:   - {_eur(s['total_retenciones'])}")
    print(f"{'─' * 65}")
    resultado = s["resultado"]
    symbol = "💸" if resultado > 0 else "💚"
    print(f"  {symbol}  RESULTADO ({s['resultado_label']}):      {_eur(abs(resultado))}")
    print(f"{'=' * 65}\n")

    # Advertencias
    if calc.engine.warnings:
        print(f"⚠️  ADVERTENCIAS ({len(calc.engine.warnings)}):")
        for w in calc.engine.warnings:
            print(f"  • {w}")
        print()


# ─── Informe Excel ────────────────────────────────────────────────────────────


def generate_excel_report(calc: IRPFCalculator, output_path: str):
    """Genera un fichero Excel con el informe completo."""
    if not HAS_OPENPYXL:
        print("  ⚠️  openpyxl no está instalado. No se puede generar el Excel.")
        return

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # eliminar hoja por defecto

    _sheet_resumen(wb, calc)
    _sheet_gpp(wb, calc.engine)
    _sheet_dividendos(wb, calc.engine)
    _sheet_inventario(wb, calc.engine)
    if calc.engine.warnings:
        _sheet_advertencias(wb, calc.engine)

    wb.save(output_path)
    print(f"  ✅  Informe Excel guardado en: {output_path}")


def _header_style(ws, row, cols, text, bg_color=COLOR_HEADER):
    """Aplica estilo de cabecera a un rango de celdas."""
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = PatternFill("solid", fgColor=bg_color)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(horizontal="center")
    ws.cell(row=row, column=1).value = text


def _col_headers(ws, row, headers, bg_color=COLOR_SUBHEAD):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.fill = PatternFill("solid", fgColor=bg_color)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")


def _auto_width(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value or "")))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


def _sheet_resumen(wb, calc: IRPFCalculator):
    ws = wb.create_sheet("Resumen IRPF")
    s = calc.summary()
    year = s["año_fiscal"]

    _header_style(ws, 1, 3, f"DECLARACIÓN DE LA RENTA {year} — Resumen IRPF")
    ws.merge_cells("A1:C1")

    sections = [
        (
            "RENDIMIENTOS DEL TRABAJO",
            COLOR_SUBHEAD,
            [
                ("Rendimiento íntegro", s["rendimiento_trabajo_bruto"]),
                ("Cotización SS trabajador", s["ss_trabajador"]),
                ("Base imponible general", s["base_general"]),
                ("Cuota íntegra (trabajo)", s["cuota_general"]),
            ],
        ),
        (
            "GANANCIAS Y PÉRDIDAS PATRIMONIALES",
            COLOR_SUBHEAD,
            [
                ("Ganancias brutas", s["ganancias_patrimoniales_bruto"]),
                ("Pérdidas brutas", s["perdidas_patrimoniales_bruto"]),
                ("GPP NETO", s["gpp_neto"]),
            ],
        ),
        (
            "RENDIMIENTOS DEL CAPITAL MOBILIARIO (Dividendos)",
            COLOR_SUBHEAD,
            [
                ("Dividendos brutos", s["dividendos_brutos"]),
                ("Retención extranjera", s["retencion_extranjera_dividendos"]),
                ("Retención española", s["retencion_espania_dividendos"]),
            ],
        ),
        (
            "BASE DEL AHORRO Y CUOTA",
            COLOR_SUBHEAD,
            [
                ("Base imponible del ahorro", s["base_ahorro"]),
                ("Cuota íntegra (ahorro)", s["cuota_ahorro"]),
                ("Deducción doble imposición", s["deduccion_doble_imposicion"]),
            ],
        ),
        (
            "RESULTADO DECLARACIÓN",
            COLOR_TOTAL,
            [
                ("Cuota íntegra total", s["cuota_integra"]),
                ("Cuota líquida", s["cuota_liquida"]),
                ("Total retenciones pagadas", s["total_retenciones"]),
                (f"RESULTADO ({s['resultado_label']})", abs(s["resultado"])),
            ],
        ),
    ]

    row = 2
    for section_title, color, items in sections:
        # Cabecera de sección
        ws.cell(row=row, column=1, value=section_title)
        ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=color)
        ws.cell(row=row, column=1).font = Font(bold=True, color="FFFFFF")
        ws.merge_cells(f"A{row}:C{row}")
        row += 1

        for label, value in items:
            ws.cell(row=row, column=1, value=label).font = Font(bold=False)
            cell_val = ws.cell(row=row, column=3, value=value)
            cell_val.number_format = '#,##0.00 "€"'
            # Color según ganancia/pérdida
            if value > 0 and "pérdida" in label.lower():
                cell_val.fill = PatternFill("solid", fgColor=COLOR_LOSS)
            elif value > 0 and any(k in label.lower() for k in ["ganancia", "resultado", "base", "cuota"]):
                cell_val.fill = PatternFill("solid", fgColor=COLOR_GAIN)
            row += 1

        row += 1  # línea en blanco entre secciones

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 5
    ws.column_dimensions["C"].width = 18


def _sheet_gpp(wb, engine: TaxEngine):
    ws = wb.create_sheet("GPP Detalle")
    year = engine.tax_year
    events = [ev for ev in engine.tax_events if ev.date.year == year]

    _header_style(ws, 1, 9, f"Ganancias y Pérdidas Patrimoniales {year}")
    ws.merge_cells("A1:I1")

    headers = [
        "Fecha venta",
        "Plataforma",
        "Activo",
        "Tipo",
        "Cantidad",
        "Fecha compra",
        "Coste adq. (€)",
        "Ingresos (€)",
        "Comisión (€)",
        "G/P (€)",
    ]
    _col_headers(ws, 2, headers)

    total_gpp = 0.0
    for i, ev in enumerate(sorted(events, key=lambda e: e.date), start=3):
        gain = ev.gain_loss_eur
        total_gpp += gain
        fill_color = COLOR_GAIN if gain >= 0 else COLOR_LOSS

        ws.cell(row=i, column=1, value=ev.date.strftime("%d/%m/%Y")).alignment = Alignment(horizontal="center")
        ws.cell(row=i, column=2, value=ev.platform)
        ws.cell(row=i, column=3, value=ev.asset)
        ws.cell(row=i, column=4, value=ev.asset_type)
        ws.cell(row=i, column=5, value=round(ev.quantity_sold, 8)).number_format = "#,##0.########"
        ws.cell(row=i, column=6, value=ev.acquisition_date.strftime("%d/%m/%Y")).alignment = Alignment(
            horizontal="center"
        )
        ws.cell(row=i, column=7, value=round(ev.acquisition_cost_eur, 2)).number_format = '#,##0.00 "€"'
        ws.cell(row=i, column=8, value=round(ev.sale_proceeds_eur, 2)).number_format = '#,##0.00 "€"'
        ws.cell(row=i, column=9, value=round(ev.fee_eur, 2)).number_format = '#,##0.00 "€"'
        cell_gp = ws.cell(row=i, column=10, value=round(gain, 2))
        cell_gp.number_format = '#,##0.00 "€"'
        cell_gp.fill = PatternFill("solid", fgColor=fill_color)

    # Fila total
    total_row = len(events) + 3
    ws.cell(row=total_row, column=9, value="TOTAL GPP").font = Font(bold=True)
    cell_total = ws.cell(row=total_row, column=10, value=round(total_gpp, 2))
    cell_total.number_format = '#,##0.00 "€"'
    cell_total.font = Font(bold=True)
    cell_total.fill = PatternFill("solid", fgColor=COLOR_TOTAL)

    _auto_width(ws)


def _sheet_dividendos(wb, engine: TaxEngine):
    ws = wb.create_sheet("Dividendos")
    year = engine.tax_year
    divs = engine.dividends

    _header_style(ws, 1, 7, f"Dividendos {year}")
    ws.merge_cells("A1:G1")

    headers = ["Fecha", "Plataforma", "Activo", "ISIN", "Bruto (€)", "Ret. extran. (€)", "Ret. España (€)"]
    _col_headers(ws, 2, headers)

    for i, d in enumerate(sorted(divs, key=lambda x: x.date), start=3):
        ws.cell(row=i, column=1, value=d.date.strftime("%d/%m/%Y")).alignment = Alignment(horizontal="center")
        ws.cell(row=i, column=2, value=d.platform)
        ws.cell(row=i, column=3, value=d.asset)
        ws.cell(row=i, column=4, value=d.isin)
        ws.cell(row=i, column=5, value=round(d.gross_eur, 2)).number_format = '#,##0.00 "€"'
        ws.cell(row=i, column=6, value=round(d.withholding_foreign_eur, 2)).number_format = '#,##0.00 "€"'
        ws.cell(row=i, column=7, value=round(d.withholding_spain_eur, 2)).number_format = '#,##0.00 "€"'

    if not divs:
        ws.cell(row=3, column=1, value="Sin dividendos en el período").font = Font(italic=True, color="808080")

    _auto_width(ws)


def _sheet_inventario(wb, engine: TaxEngine):
    ws = wb.create_sheet("Inventario Final")
    _header_style(ws, 1, 5, f"Inventario de activos al 31/12/{engine.tax_year}")
    ws.merge_cells("A1:E1")

    headers = ["Activo", "Cantidad", "Fecha adquisición", "Coste/unidad (€)", "Coste total (€)"]
    _col_headers(ws, 2, headers)

    row = 3
    for asset, lots in sorted(engine.inventory.items()):
        for lot in lots:
            if lot.quantity < 1e-10:
                continue
            ws.cell(row=row, column=1, value=asset)
            ws.cell(row=row, column=2, value=round(lot.quantity, 8)).number_format = "#,##0.########"
            ws.cell(row=row, column=3, value=lot.date.strftime("%d/%m/%Y")).alignment = Alignment(horizontal="center")
            ws.cell(row=row, column=4, value=round(lot.cost_per_unit_eur, 4)).number_format = '#,##0.0000 "€"'
            ws.cell(row=row, column=5, value=round(lot.total_cost_eur, 2)).number_format = '#,##0.00 "€"'
            row += 1

    if row == 3:
        ws.cell(row=3, column=1, value="Sin inventario pendiente").font = Font(italic=True, color="808080")

    _auto_width(ws)


def _sheet_advertencias(wb, engine: TaxEngine):
    ws = wb.create_sheet("⚠️ Advertencias")
    _header_style(ws, 1, 2, "Advertencias — revisar manualmente", bg_color="C00000")
    ws.merge_cells("A1:B1")
    ws.cell(row=1, column=1).font = Font(bold=True, color="FFFFFF", size=12)

    for i, w in enumerate(engine.warnings, start=2):
        ws.cell(row=i, column=1, value=f"•  {w}")
        ws.cell(row=i, column=1).fill = PatternFill("solid", fgColor="FFE0E0")

    ws.column_dimensions["A"].width = 90

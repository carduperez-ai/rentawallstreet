import pandas as pd
from typing import List
from src.domain.crypto_entities import TaxEvent
from src.domain.fiscal_entities import Dividend
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


def generate_audit_excel(tax_events: List[TaxEvent], dividends: List[Dividend], output_path: str):
    """
    Genera un informe Excel altamente estructurado para auditoría fiscal (Libro Diario de FIFO).
    """
    try:
        # 1. Preparar datos de Transmisiones (GPP) con estructura de "Libro Diario"
        rows = []
        # Agrupar eventos por fecha y activo para identificar la "Operación de Venta" original
        # Ya que un solo trade de venta puede descomponerse en varios eventos TaxEvent (lotes FIFO)

        # Ordenamos por fecha y activo
        sorted_events = sorted(tax_events, key=lambda x: (x.date, x.asset))

        current_sale_id = 0
        last_key = None

        for ev in sorted_events:
            key = (ev.date, ev.asset, ev.platform)
            if key != last_key:
                current_sale_id += 1
                last_key = key

            rows.append(
                {
                    "ID Operación": f"VENTA-{current_sale_id:03d}",
                    "Fecha Venta": ev.date.strftime("%d/%m/%Y %H:%M"),
                    "Activo": ev.asset,
                    "Qty Vendida (Tramo)": ev.quantity_sold,
                    "Valor Transmisión (€)": ev.sale_proceeds_eur,
                    "Gastos Venta (€)": ev.sell_fee_eur,
                    "|": "-->",
                    "Fecha Adquisición": ev.acquisition_date.strftime("%d/%m/%Y"),
                    "Coste Adquisición (€)": ev.acquisition_cost_eur,
                    "Gastos Compra (€)": ev.buy_fee_eur,
                    "Resultado Tramo (€)": ev.gain_loss_eur,
                    "Plataforma": ev.platform,
                    "Detalle / Auditoría": ev.notes,
                }
            )

        df_gpp = pd.DataFrame(rows)

        # 2. Preparar datos de Rendimientos (RCM) - SOLO DIVIDENDOS REALES
        rcm_rows = []
        # Filtramos para que solo aparezcan dividendos en la tabla [0029]
        real_dividends = [d for d in dividends if getattr(d, "type", "dividend") == "dividend"]

        for d in real_dividends:
            rcm_rows.append(
                {
                    "Fecha": d.date.strftime("%d/%m/%Y"),
                    "Activo": d.asset,
                    "Concepto": "Dividendo / Rendimiento",
                    "Importe Bruto (€)": d.gross_eur,
                    "Retención Origen (€)": d.withholding_foreign_eur,
                    "Retención España (€)": d.withholding_spain_eur,
                    "Plataforma": d.platform,
                }
            )
        df_rcm = pd.DataFrame(rcm_rows)

        # 2b. Preparar datos de Gastos Deducibles [0035]
        fee_rows = []
        real_fees = [d for d in dividends if getattr(d, "type", "dividend") == "fee"]
        for f in real_fees:
            fee_rows.append(
                {
                    "Fecha": f.date.strftime("%d/%m/%Y"),
                    "Concepto": f.asset,
                    "Importe Deducible (€)": f.custody_fee_eur,
                    "Plataforma": f.platform,
                }
            )
        df_fees = pd.DataFrame(fee_rows)

        # 3. Escribir a Excel con formato avanzado
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            if not df_gpp.empty:
                df_gpp.to_excel(writer, sheet_name="Libro Diario FIFO", index=False)
            if not df_rcm.empty:
                df_rcm.to_excel(writer, sheet_name="Rendimientos (RCM) [0029]", index=False)
            if not df_fees.empty:
                df_fees.to_excel(writer, sheet_name="Gastos Deducibles [0035]", index=False)

            # --- ESTILOS ---
            header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)
            border = Border(
                left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin")
            )

            # Aplicar formato a la hoja principal
            if "Libro Diario FIFO" in writer.sheets:
                ws = writer.sheets["Libro Diario FIFO"]
                # Congelar paneles
                ws.freeze_panes = "A2"

                # Formatear cabecera
                for cell in ws[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center")

                # Ajustar columnas y agrupar visualmente por ID Operación
                for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
                    # Alternar color por ID Operación para facilitar lectura
                    op_id = ws.cell(row=i, column=1).value

                    fill_color = "F2F2F2" if (int(op_id.split("-")[1]) % 2 == 0) else "FFFFFF"
                    row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

                    for cell in row:
                        cell.fill = row_fill
                        cell.border = border
                        if isinstance(cell.value, (int, float)):
                            cell.number_format = "#,##0.00"

                # Ajustar anchos
                col_widths = {
                    "A": 15,
                    "B": 20,
                    "C": 10,
                    "D": 15,
                    "E": 20,
                    "F": 15,
                    "G": 5,
                    "H": 18,
                    "I": 20,
                    "J": 15,
                    "K": 18,
                    "L": 15,
                    "M": 40,
                }
                for col, width in col_widths.items():
                    ws.column_dimensions[col].width = width

        return True
    except Exception as e:
        print(f"Error generando Excel: {e}")
        return False

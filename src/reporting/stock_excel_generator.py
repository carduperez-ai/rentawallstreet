import pandas as pd
from typing import List
from src.domain.stock_entities import Trade, TaxEvent
from src.domain.stock_fiscal_entities import Dividend
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

def generate_stock_audit_excel(tax_events: List[TaxEvent], dividends: List[Dividend], output_path: str):
    """
    Genera un informe Excel para auditoría fiscal de Acciones y ETFs (Libro Diario de FIFO).
    """
    try:
        rows = []
        sorted_events = sorted(tax_events, key=lambda x: (x.date, x.asset))
        
        current_sale_id = 0
        last_key = None
        
        for ev in sorted_events:
            key = (ev.date, ev.asset, ev.platform)
            if key != last_key:
                current_sale_id += 1
                last_key = key
            
            rows.append({
                "ID Operación": f"VENTA-STOCK-{current_sale_id:03d}",
                "Fecha Venta": ev.date.strftime('%d/%m/%Y %H:%M'),
                "Activo": ev.asset,
                "Qty Vendida": ev.quantity_sold,
                "Valor Transmisión (€)": ev.sale_proceeds_eur,
                "Gastos Venta (€)": ev.sell_fee_eur,
                "|": "-->",
                "Fecha Adquisición": ev.acquisition_date.strftime('%d/%m/%Y'),
                "Coste Adquisición (€)": ev.acquisition_cost_eur,
                "Gastos Compra (€)": ev.buy_fee_eur,
                "Resultado Tramo (€)": ev.gain_loss_eur,
                "Plataforma": ev.platform,
                "Detalle": ev.notes
            })
        
        df_gpp = pd.DataFrame(rows)

        rcm_rows = []
        for d in dividends:
            rcm_rows.append({
                "Fecha": d.date.strftime('%d/%m/%Y'),
                "Activo": d.asset,
                "ISIN": d.isin,
                "Concepto": "Dividendo / Cupón",
                "Importe Bruto (€)": d.gross_eur,
                "Retención Origen (€)": d.withholding_foreign_eur,
                "Retención España (€)": d.withholding_spain_eur,
                "Plataforma": d.platform
            })
        df_rcm = pd.DataFrame(rcm_rows)

        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            if not df_gpp.empty:
                df_gpp.to_excel(writer, sheet_name="Libro Diario FIFO Acciones", index=False)
            if not df_rcm.empty:
                df_rcm.to_excel(writer, sheet_name="Rendimientos Bolsa (RCM)", index=False)
            
            # --- ESTILOS (Verde para Bolsa) ---
            header_fill = PatternFill(start_color="375623", end_color="375623", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)
            border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
            
            if "Libro Diario FIFO Acciones" in writer.sheets:
                ws = writer.sheets["Libro Diario FIFO Acciones"]
                ws.freeze_panes = "A2"
                for cell in ws[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center")
                
                for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
                    fill_color = "F2F9F2" if (i % 2 == 0) else "FFFFFF"
                    row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
                    for cell in row:
                        cell.fill = row_fill
                        cell.border = border
                        if isinstance(cell.value, (int, float)):
                            cell.number_format = '#,##0.00'
                
                col_widths = {"A": 20, "B": 20, "C": 15, "D": 15, "E": 20, "F": 15, "G": 5, "H": 18, "I": 20, "J": 15, "K": 18, "L": 15, "M": 40}
                for col, width in col_widths.items():
                    ws.column_dimensions[col].width = width

        return True
    except Exception as e:
        print(f"Error generando Excel Acciones: {e}")
        return False
